"""Export des alertes en CSV (sans dépendance) et Excel (openpyxl si présent)."""
import csv
import io

from ..models import Alert
from .alerts import list_alerts

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    OPENPYXL_AVAILABLE = True
except Exception:  # pragma: no cover
    OPENPYXL_AVAILABLE = False

COLUMNS = [
    ("id", "ID"), ("time_full", "Date/heure"), ("type", "Type"),
    ("urgency", "Urgence"), ("status", "Statut"), ("neighborhood", "Quartier"),
    ("lat", "Latitude"), ("lng", "Longitude"),
    ("reporter_name", "Citoyen"), ("reporter_phone", "Téléphone"),
    ("agent", "Agent"), ("distance_m", "Distance (m)"), ("description", "Description"),
]


def _rows(filters):
    # Pas de pagination : on exporte tout ce qui correspond aux filtres.
    result = list_alerts(filters=filters, page=1, page_size=100000)
    for a in result["items"]:
        yield {
            "id": a["id"],
            "time_full": (a.get("created_at") or "").replace("T", " ")[:19],
            "type": a.get("type"),
            "urgency": a.get("urgency"),
            "status": a.get("status"),
            "neighborhood": a.get("neighborhood") or "",
            "lat": a.get("lat"), "lng": a.get("lng"),
            "reporter_name": a.get("reporter_name") or "",
            "reporter_phone": a.get("reporter_phone") or "",
            "agent": (a.get("assigned_agent") or {}).get("name", "") if a.get("assigned_agent") else "",
            "distance_m": a.get("distance_m"),
            "description": a.get("description") or "",
        }


def to_csv(filters=None):
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow([label for _, label in COLUMNS])
    for row in _rows(filters or {}):
        writer.writerow([row[key] for key, _ in COLUMNS])
    # BOM pour qu'Excel ouvre l'UTF-8 correctement.
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(filters=None):
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError("openpyxl n'est pas installé (pip install openpyxl).")
    wb = Workbook()
    ws = wb.active
    ws.title = "Alertes"
    header_fill = PatternFill("solid", fgColor="0B1220")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append([label for _, label in COLUMNS])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in _rows(filters or {}):
        ws.append([row[key] for key, _ in COLUMNS])
    # Largeurs de colonnes approximatives.
    widths = [6, 20, 12, 10, 10, 14, 11, 11, 18, 16, 16, 12, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
